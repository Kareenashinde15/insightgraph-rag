from typing import List
from ingestion.parsers.base import BaseParser, ParsedDocument, ParsedChunkDraft

class DocxParser(BaseParser):
    def parse(self, file_path: str, filename: str) -> ParsedDocument:
        drafts: List[ParsedChunkDraft] = []
        raw_text_parts: List[str] = []

        try:
            import docx
            doc = docx.Document(file_path)
            current_section = "Overview"
            current_page = 1
            para_count = 0

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                raw_text_parts.append(text)
                para_count += 1
                if para_count % 8 == 0:
                    current_page += 1

                if p.style and p.style.name and p.style.name.startswith("Heading"):
                    current_section = text
                    continue

                drafts.append(
                    ParsedChunkDraft(
                        page_number=current_page,
                        section=current_section,
                        paragraph=text,
                        text=text,
                    )
                )

            # Also parse tables if any
            for table in doc.tables:
                table_lines = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    table_lines.append(" | ".join(row_cells))
                if table_lines:
                    table_text = "\n".join(table_lines)
                    raw_text_parts.append(table_text)
                    drafts.append(
                        ParsedChunkDraft(
                            page_number=current_page,
                            section=f"{current_section} - Table",
                            paragraph=table_text,
                            text=table_text,
                        )
                    )

        except Exception as e:
            drafts.append(
                ParsedChunkDraft(
                    page_number=1,
                    section="Error/Fallback",
                    paragraph=f"Could not parse DOCX: {str(e)}",
                    text=f"Could not parse DOCX: {str(e)}",
                )
            )

        full_raw = "\n\n".join(raw_text_parts)
        return ParsedDocument(
            filename=filename,
            file_type="docx",
            raw_text=full_raw,
            pages=drafts,
            metadata={"paragraphs_count": len(drafts)},
        )
